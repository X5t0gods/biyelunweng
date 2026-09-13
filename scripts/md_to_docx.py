"""把 Markdown 开题报告转换为符合高校格式要求的 Word 文档。

运行：python scripts/md_to_docx.py [输入md] [输出docx]
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DEFAULT_INPUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                             'docs', '开题报告.md')
DEFAULT_OUTPUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                              'docs', '开题报告.docx')

FONT_SONG = '宋体'
FONT_HEI = '黑体'
FONT_KAI = '楷体'


def set_run_font(run, font_name, size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_paragraph_text(doc, text, font=FONT_SONG, size=12, bold=False,
                       align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=True,
                       space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if first_line_indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_run_font(run, font, size, bold)
    return p


def parse_table_row(line):
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return cells


def is_separator(line):
    return bool(re.match(r'^\s*\|[\s:\-\|]+\|\s*$', line))


def clean_inline(text):
    """去除 markdown 行内标记。"""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'\[(.+?)\]\(.*?\)', r'\1', text)
    return text.strip()


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    style = doc.styles['Normal']
    style.font.name = FONT_SONG
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_SONG)

    i = 0
    in_code = False
    code_buf = []
    first_title_done = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith('```'):
            if in_code:
                if code_buf:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run('\n'.join(code_buf))
                    set_run_font(run, 'Consolas', 10)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.match(r'^-{3,}$', stripped):
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and is_separator(lines[i + 1]):
            header = parse_table_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for idx, text in enumerate(header):
                hdr[idx].text = ''
                p = hdr[idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(clean_inline(text))
                set_run_font(run, FONT_HEI, 10.5, bold=True)
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    text = clean_inline(row[idx]) if idx < len(row) else ''
                    cells[idx].text = ''
                    p = cells[idx].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 12 else WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(text)
                    set_run_font(run, FONT_SONG, 10.5)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = clean_inline(m.group(2))
            if level == 1 and not first_title_done:
                first_title_done = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
                run = p.add_run(text)
                set_run_font(run, FONT_HEI, 18, bold=True)
            elif first_title_done and level == 1 and text.startswith('开题报告'):
                # 副标题（文档中部出现的一级标题）
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                run = p.add_run(text)
                set_run_font(run, FONT_HEI, 16, bold=True)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.5
                size = {1: 16, 2: 14, 3: 12.5, 4: 12}[level]
                run = p.add_run(text)
                set_run_font(run, FONT_HEI, size, bold=True)
            i += 1
            continue

        # 列表
        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            add_paragraph_text(doc, '· ' + clean_inline(m.group(1)), first_line_indent=False)
            i += 1
            continue
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            add_paragraph_text(doc, f'{m.group(1)}. {clean_inline(m.group(2))}',
                               first_line_indent=False)
            i += 1
            continue
        m = re.match(r'^>\s*(.*)$', stripped)
        if m:
            add_paragraph_text(doc, clean_inline(m.group(1)), font=FONT_KAI,
                               first_line_indent=True)
            i += 1
            continue

        # 普通段落（合并多行）
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith('#') or nxt.startswith('|') or nxt.startswith('```')
                    or re.match(r'^[-*]\s+', nxt) or re.match(r'^\d+\.\s+', nxt)
                    or re.match(r'^-{3,}$', nxt) or nxt.startswith('>')):
                break
            buf.append(nxt)
            i += 1
        add_paragraph_text(doc, clean_inline(' '.join(buf)))

    doc.save(docx_path)
    print(f'已生成 Word 文档：{docx_path}')


def main():
    md_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_INPUT
    docx_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
    convert(md_path, docx_path)


if __name__ == '__main__':
    main()
